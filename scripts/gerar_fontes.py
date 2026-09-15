import csv
import json
import random
import uuid
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook

RAIZ = Path(__file__).resolve().parents[1]
RAW = RAIZ / "data" / "raw"

random.seed(20260917)

AREAS = [
    (1101, "1ª Vara Cível de Campo Grande", "Judicial", "Campo Grande", "1ª Circunscrição", date(2026, 2, 10)),
    (1102, "Juizado Especial Cível de Campo Grande", "JUDICIAL", "CAMPO GRANDE", "1ª Circunscrição", date(2026, 2, 10)),
    (1201, "2ª Vara Criminal de Dourados", "Judicial", "Comarca de Dourados", "2ª Circunscrição", "03/03/2026"),
    (1202, "Vara de Família e Sucessões de Dourados", "judicial ", " dourados", "2ª Circunscrição", date(2026, 3, 3)),
    (1301, "Cartório da 1ª Vara de Três Lagoas", "Judicial", "Três Lagoas", "3ª Circunscrição", date(2026, 4, 22)),
    (1302, "Vara da Infância e Juventude de Três Lagoas", "Judicial", "TRÊS LAGOAS ", "3ª Circunscrição", date(2026, 4, 22)),
    (1401, "Vara Única de Rio Brilhante", "JUDICIAL", "Rio Brilhante", "2ª Circunscrição", "15/01/2026"),
    (1501, "1ª Vara de Corumbá", "Judicial", "Comarca de Corumbá", "4ª Circunscrição", date(2026, 1, 28)),
    (1502, "Juizado Especial Adjunto de Corumbá", "Judicial", "corumbá", None, date(2026, 1, 28)),
    (1601, "2ª Vara de Ponta Porã", "Judicial", "Ponta Porã", "2ª Circunscrição", date(2026, 2, 17)),
    (1701, "Vara Única de Aquidauana", "judicial", "Comarca de Aquidauana", "4ª Circunscrição", "09/02/2026"),
    (9001, "Secretaria de Tecnologia da Informação", "Adm.", None, None, date(2026, 5, 5)),
    (9002, "Secretaria de Material e Patrimônio", "ADMINISTRATIVA", None, None, date(2026, 5, 5)),
]

AREA_DUPLICADA = (1301, "CARTORIO DA 1A VARA DE TRES LAGOAS", "JUDICIAL", "TRES LAGOAS", None, date(2024, 11, 8))

AREA_EXTINTA = 1205

VIZINHAS = {
    1101: [1102, 9002],
    1102: [1101],
    1201: [1202],
    1202: [1201],
    1205: [1201, 1202],
    1301: [1302],
    1302: [1301],
    1401: [1201, 1202],
    1501: [1502],
    1502: [1501],
    1601: [1201],
    1701: [1501],
    9001: [9002, 1101, 1102],
    9002: [9001],
}

PERFIL = {
    1101: {2025: (0.93, 0.03), 2026: (0.95, 0.02)},
    1102: {2025: (0.89, 0.05), 2026: (0.91, 0.04)},
    1201: {2025: (0.86, 0.06), 2026: (0.88, 0.05)},
    1202: {2025: (0.90, 0.03), 2026: (0.90, 0.03)},
    1205: {2025: (0.00, 0.80), 2026: (0.00, 0.75)},
    1301: {2025: (0.84, 0.07), 2026: (0.86, 0.06)},
    1302: {2025: (0.80, 0.09), 2026: (0.83, 0.08)},
    1401: {2025: (0.61, 0.32), 2026: (0.52, 0.41)},
    1501: {2025: (0.79, 0.05), 2026: (0.81, 0.05)},
    1502: {2025: (0.70, 0.06), 2026: (0.74, 0.06)},
    1601: {2025: (0.87, 0.04), 2026: (0.89, 0.04)},
    1701: {2025: (0.76, 0.07), 2026: (0.72, 0.08)},
    9001: {2025: (0.81, 0.11), 2026: (0.84, 0.10)},
    9002: {2025: (0.96, 0.01), 2026: (0.96, 0.01)},
}

CAPTURA = {
    1502: (0.0, 0.70, 0.30),
    1701: (0.0, 0.72, 0.28),
    9001: (0.72, 0.24, 0.04),
}
CAPTURA_PADRAO = (0.55, 0.36, 0.09)

QUANTIDADE_BENS = {
    1101: 128, 1102: 96, 1201: 117, 1202: 84, 1205: 6, 1301: 102, 1302: 71,
    1401: 88, 1501: 93, 1502: 58, 1601: 79, 1701: 66, 9001: 243, 9002: 162,
}

ITENS = [
    ("Cadeira giratória com braços", 280, 950),
    ("Mesa de escritório em L", 450, 1600),
    ("Armário de aço 2 portas", 600, 1900),
    ("Gaveteiro volante", 220, 700),
    ("Longarina 3 lugares", 380, 1100),
    ("Condicionador de ar split", 1900, 4600),
    ("Estação de trabalho", 2800, 5200),
    ("Notebook", 3200, 6800),
    ("Monitor LED 24 polegadas", 650, 1400),
    ("Impressora multifuncional", 1300, 4200),
    ("Nobreak 1500VA", 450, 1300),
    ("Scanner de mesa", 900, 2600),
]
PESOS_JUDICIAL = [14, 10, 8, 8, 5, 4, 10, 4, 10, 4, 5, 3]
PESOS_TI = [4, 3, 2, 2, 1, 2, 18, 22, 22, 6, 12, 6]

ESPACOS = ["Gabinete", "Cartório", "Sala de audiências", "Arquivo", "Recepção", "Almoxarifado"] + [
    f"Sala {n}" for n in range(101, 121)
]

JANELAS = {
    2025: {"inicio": date(2025, 10, 6), "dias": 40},
    2026: {"inicio": date(2026, 8, 3), "dias": 26},
}
CORTE_LOTE_2026 = date(2026, 8, 14)


def novo_uuid() -> str:
    return str(uuid.UUID(int=random.getrandbits(128), version=4))


def valor_br(valor: float) -> str:
    inteiro, centavos = f"{valor:.2f}".split(".")
    return f"{int(inteiro):,}".replace(",", ".") + "," + centavos


def formatar_numero_carga(numero: int) -> str:
    sorteio = random.random()
    if sorteio < 0.05:
        return f"{numero:08d}"
    if sorteio < 0.10:
        texto = str(numero)
        return f"{texto[:3]}.{texto[3:]}"
    return str(numero)


def novo_bem(numeros, area: int) -> dict:
    pesos = PESOS_TI if area == 9001 else PESOS_JUDICIAL
    descricao, minimo, maximo = random.choices(ITENS, weights=pesos)[0]
    return {
        "numero": next(numeros),
        "descricao": descricao,
        "valor": round(random.uniform(minimo, maximo), 2),
        "area": area,
    }


def gerar_bens():
    numeros = iter(random.sample(range(100000, 999999), 8000))
    bens_2025 = [novo_bem(numeros, area) for area, qtd in QUANTIDADE_BENS.items() for _ in range(qtd)]

    bens_2026 = []
    for bem in bens_2025:
        if random.random() < 0.03:
            continue
        atual = dict(bem, valor=round(bem["valor"] * random.uniform(0.82, 0.92), 2))
        if atual["area"] != AREA_EXTINTA and random.random() < 0.02:
            atual["area"] = random.choice(VIZINHAS[atual["area"]])
        bens_2026.append(atual)
    for area, qtd in QUANTIDADE_BENS.items():
        if area == AREA_EXTINTA:
            continue
        bens_2026.extend(novo_bem(numeros, area) for _ in range(round(qtd * 0.06)))

    for bens in (bens_2025, bens_2026):
        for bem in random.sample(bens, 11):
            bem["baixado"] = True

    return {2025: bens_2025, 2026: bens_2026}


def linhas_carga(bens_por_exercicio) -> list[list[str]]:
    linhas = []
    for exercicio, bens in bens_por_exercicio.items():
        bloco = []
        for bem in bens:
            valor = valor_br(bem["valor"]) if random.random() < 0.9 else f"{bem['valor']:.2f}"
            situacao = "BAIXADO" if bem.get("baixado") else random.choices(["ATIVO", "Ativo"], weights=[85, 15])[0]
            bloco.append([
                str(exercicio),
                formatar_numero_carga(bem["numero"]),
                bem["descricao"],
                str(bem["area"]),
                valor,
                situacao,
            ])
        random.shuffle(bloco)
        linhas.extend(bloco)

    for _ in range(4):
        posicao = random.randrange(len(linhas))
        linhas.insert(random.randrange(len(linhas)), list(linhas[posicao]))
    return linhas


def sortear_captura(area: int) -> str:
    leitor, celular, manual = CAPTURA.get(area, CAPTURA_PADRAO)
    return random.choices(["leitor", "celular", "manual"], weights=[leitor, celular, manual])[0]


def origem_informada(captura: str, exercicio: int):
    if exercicio == 2025:
        return {"leitor": "manual", "celular": "scanner", "manual": "manual"}[captura]
    if random.random() < 0.006:
        return None
    if captura == "leitor":
        return "leitor_codigo_barras"
    if captura == "celular":
        return random.choices(["celular", "mobile_celular"], weights=[70, 30])[0]
    return random.choices(["manual", "mobile_manual"], weights=[75, 25])[0]


def numero_lido(numero: int, captura: str) -> str:
    if captura == "leitor":
        return f"{numero:08d}"
    if captura == "celular":
        return str(numero)
    sorteio = random.random()
    if sorteio < 0.35:
        return f" {numero} "
    if sorteio < 0.60:
        texto = str(numero)
        return f"{texto[:3]}.{texto[3:]}"
    return str(numero)


def momento_coleta(exercicio: int, area: int, captura: str) -> str:
    janela = JANELAS[exercicio]
    dia_area = janela["inicio"] + timedelta(days=(area * 7) % janela["dias"])
    quando = datetime.combine(dia_area, datetime.min.time()) + timedelta(
        days=random.randint(0, 3), hours=random.randint(8, 17), minutes=random.randint(0, 59), seconds=random.randint(0, 59)
    )
    if captura == "manual" and random.random() < 0.7:
        return quando.strftime("%d/%m/%Y %H:%M")
    return quando.strftime("%Y-%m-%dT%H:%M:%S")


def evento(exercicio: int, numero: int, area: int, captura: str, espaco: str, coletor: str) -> dict:
    return {
        "id_leitura": novo_uuid(),
        "exercicio": exercicio,
        "numero_patrimonio": numero_lido(numero, captura),
        "local": {"codigo_area": area, "espaco_fisico": espaco},
        "origem_leitura": origem_informada(captura, exercicio),
        "coletor": coletor,
        "coletado_em": momento_coleta(exercicio, area, captura),
    }


def gerar_leituras(exercicio: int, bens: list[dict]) -> list[dict]:
    coletores = {area: [f"srv{random.randint(1000, 9999)}" for _ in range(3)] for area in QUANTIDADE_BENS}
    eventos = []
    for bem in bens:
        if bem.get("baixado"):
            continue
        local, outra = PERFIL[bem["area"]][exercicio]
        sorteio = random.random()
        if sorteio < local:
            destino = bem["area"]
        elif sorteio < local + outra:
            destino = random.choice(VIZINHAS[bem["area"]])
        else:
            continue

        leituras_do_bem = [destino]
        if random.random() < 0.05:
            opcoes = [a for a in [bem["area"]] + VIZINHAS[bem["area"]] if a != AREA_EXTINTA]
            primeira = random.choice(opcoes)
            leituras_do_bem.insert(0, primeira)

        horarios = []
        for area_lida in leituras_do_bem:
            captura = sortear_captura(area_lida)
            item = evento(exercicio, bem["numero"], area_lida, captura, random.choice(ESPACOS), random.choice(coletores[area_lida]))
            horarios.append(item)
        if len(horarios) == 2:
            primeiro, segundo = horarios
            if segundo["coletado_em"] <= primeiro["coletado_em"] or "/" in segundo["coletado_em"] or "/" in primeiro["coletado_em"]:
                base = datetime.strptime(momento_coleta(exercicio, destino, "celular"), "%Y-%m-%dT%H:%M:%S")
                primeiro["coletado_em"] = (base - timedelta(days=1)).strftime("%Y-%m-%dT%H:%M:%S")
                segundo["coletado_em"] = base.strftime("%Y-%m-%dT%H:%M:%S")
        eventos.extend(horarios)

    todos_numeros = {bem["numero"] for bem in bens}
    areas_validas = [a for a in QUANTIDADE_BENS if a != AREA_EXTINTA]
    for _ in range(9):
        area = random.choice(areas_validas)
        numero = random.choice(list(todos_numeros)) // 10
        eventos.append(evento(exercicio, numero, area, "manual", random.choice(ESPACOS), random.choice(coletores[area])))
    for _ in range(3):
        area = random.choice(areas_validas)
        item = evento(exercicio, 0, area, "manual", random.choice(ESPACOS), random.choice(coletores[area]))
        item["numero_patrimonio"] = ""
        eventos.append(item)

    random.shuffle(eventos)
    return eventos


def data_do_evento(item: dict) -> date:
    texto = item["coletado_em"]
    if "/" in texto:
        return datetime.strptime(texto, "%d/%m/%Y %H:%M").date()
    return datetime.strptime(texto, "%Y-%m-%dT%H:%M:%S").date()


def salvar_json(caminho: Path, exercicio: int, gerado_em: str, leituras: list[dict]) -> None:
    payload = {"sistema": "coleta-inventario", "exercicio": exercicio, "gerado_em": gerado_em, "leituras": leituras}
    caminho.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def salvar_areas(caminho: Path) -> None:
    planilha = Workbook()
    aba = planilha.active
    aba.title = "areas"
    aba.append(["codigo_area", "nome_area", "tipo_area", "comarca", "regiao", "atualizado_em"])
    linhas = list(AREAS)
    linhas.insert(9, AREA_DUPLICADA)
    for codigo, nome, tipo, comarca, regiao, atualizado in linhas:
        aba.append([codigo, nome, tipo, comarca, regiao, atualizado])
        if isinstance(atualizado, date):
            aba.cell(row=aba.max_row, column=6).number_format = "DD/MM/YYYY"
    aba.column_dimensions["B"].width = 44
    aba.column_dimensions["D"].width = 24
    aba.column_dimensions["E"].width = 18
    planilha.save(caminho)


def main() -> None:
    (RAW / "leituras").mkdir(parents=True, exist_ok=True)

    salvar_areas(RAW / "areas.xlsx")

    bens = gerar_bens()
    with open(RAW / "carga_patrimonial.csv", "w", newline="", encoding="latin-1") as arquivo:
        escritor = csv.writer(arquivo, delimiter=";", lineterminator="\r\n")
        escritor.writerow(["exercicio", "numero_patrimonio", "descricao_bem", "codigo_area_carga", "valor_atual", "situacao_bem"])
        escritor.writerows(linhas_carga(bens))

    leituras_2025 = gerar_leituras(2025, bens[2025])
    salvar_json(RAW / "leituras" / "leituras_2025.json", 2025, "2025-11-21T18:40:02", leituras_2025)

    leituras_2026 = gerar_leituras(2026, bens[2026])
    parte1 = [e for e in leituras_2026 if data_do_evento(e) <= CORTE_LOTE_2026]
    parte2 = [e for e in leituras_2026 if data_do_evento(e) > CORTE_LOTE_2026]
    parte2.extend(dict(e) for e in random.sample(parte1, 12))
    random.shuffle(parte2)
    salvar_json(RAW / "leituras" / "leituras_2026_parte1.json", 2026, "2026-08-14T19:02:11", parte1)
    salvar_json(RAW / "leituras" / "leituras_2026_parte2.json", 2026, "2026-08-28T18:55:47", parte2)

    print(f"areas.xlsx                  {len(AREAS) + 1} linhas")
    print(f"carga_patrimonial.csv       {sum(len(b) for b in bens.values()) + 4} linhas")
    print(f"leituras_2025.json          {len(leituras_2025)} leituras")
    print(f"leituras_2026_parte1.json   {len(parte1)} leituras")
    print(f"leituras_2026_parte2.json   {len(parte2)} leituras")


if __name__ == "__main__":
    main()
